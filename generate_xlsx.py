"""
Generate Company_products.xlsx sample data for ACME Corporation.
This script creates the Excel file that will be loaded by the RAG pipeline.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Product Catalog"

# Header styling
header_font = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
header_fill = PatternFill(start_color="6366F1", end_color="6366F1", fill_type="solid")
header_alignment = Alignment(horizontal="center", vertical="center")
thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

headers = ["Product ID", "Product Name", "Category", "Price (USD)", "Description", "Availability", "License Type"]
for col, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment
    cell.border = thin_border

# Product data
products = [
    ["PRD-001", "ACME Enterprise Suite - Basic", "Software Suite", 299.99,
     "Entry-level enterprise solution with core CRM, project management, and basic analytics. Includes email support and knowledge base access. Ideal for small teams up to 25 users.",
     "In Stock", "Annual Subscription"],
    ["PRD-002", "ACME Enterprise Suite - Professional", "Software Suite", 599.99,
     "Mid-tier enterprise solution with advanced CRM, workflow automation, custom dashboards, and API access. Includes priority email and chat support. Supports up to 100 users.",
     "In Stock", "Annual Subscription"],
    ["PRD-003", "ACME Enterprise Suite - Premium", "Software Suite", 1199.99,
     "Full-featured enterprise solution with AI-powered analytics, unlimited integrations, custom development support, dedicated account manager, and 24/7 phone support. Unlimited users.",
     "In Stock", "Annual Subscription"],
    ["PRD-004", "ACME CRM Platform", "CRM", 149.99,
     "Standalone CRM platform with contact management, lead tracking, sales pipeline visualization, email campaigns, and reporting. Integrates with major email providers and calendars.",
     "In Stock", "Monthly Subscription"],
    ["PRD-005", "ACME Data Analytics Dashboard", "Analytics", 249.99,
     "Powerful data visualization and analytics tool with 50+ chart types, real-time data connectors, custom KPI tracking, automated reporting, and team collaboration features.",
     "In Stock", "Monthly Subscription"],
    ["PRD-006", "ACME Cloud Security Shield", "Security", 399.99,
     "Comprehensive cloud security solution with threat detection, vulnerability scanning, compliance monitoring (GDPR, HIPAA, PCI), incident response automation, and security audit logs.",
     "In Stock", "Annual Subscription"],
    ["PRD-007", "ACME AI Implementation Toolkit", "AI/ML", 799.99,
     "Complete toolkit for deploying AI solutions including pre-trained models, custom model training, MLOps pipeline, model monitoring, and deployment automation. Includes 100 GPU hours/month.",
     "In Stock", "Annual Subscription"],
    ["PRD-008", "ACME Workflow Automator", "Automation", 199.99,
     "No-code workflow automation platform with drag-and-drop builder, 200+ pre-built templates, multi-step workflows, conditional logic, and integration with 500+ apps.",
     "In Stock", "Monthly Subscription"],
    ["PRD-009", "ACME Team Collaboration Hub", "Collaboration", 89.99,
     "Team communication and collaboration platform with channels, direct messaging, file sharing, video conferencing, task boards, and integration with popular productivity tools.",
     "In Stock", "Monthly Subscription"],
    ["PRD-010", "ACME DevOps Pipeline Pro", "DevOps", 349.99,
     "End-to-end DevOps solution with CI/CD pipelines, infrastructure-as-code, container orchestration, monitoring, and automated rollbacks. Supports Kubernetes, Docker, and major cloud providers.",
     "In Stock", "Monthly Subscription"],
    ["PRD-011", "ACME Customer Insights Engine", "Analytics", 449.99,
     "AI-powered customer analytics platform with sentiment analysis, churn prediction, customer segmentation, lifetime value modeling, and personalized recommendation engine.",
     "In Stock", "Annual Subscription"],
    ["PRD-012", "ACME Compliance Manager", "Compliance", 299.99,
     "Regulatory compliance management tool with policy tracking, audit trails, risk assessments, training management, and automated compliance reporting for SOX, GDPR, HIPAA, and PCI.",
     "Limited Stock", "Annual Subscription"],
    ["PRD-013", "ACME API Gateway", "Infrastructure", 179.99,
     "Enterprise API management platform with rate limiting, authentication, analytics, documentation generation, versioning, and developer portal. Handles up to 10M requests/month.",
     "In Stock", "Monthly Subscription"],
    ["PRD-014", "ACME Backup & Recovery Pro", "Infrastructure", 129.99,
     "Automated backup solution with incremental backups, point-in-time recovery, cross-region replication, encryption at rest, and compliance-ready retention policies.",
     "In Stock", "Monthly Subscription"],
    ["PRD-015", "ACME Smart IoT Platform", "IoT", 549.99,
     "IoT device management and analytics platform with device provisioning, real-time monitoring, edge computing support, predictive maintenance, and OTA firmware updates.",
     "Pre-Order", "Annual Subscription"],
]

for row_idx, product in enumerate(products, 2):
    for col_idx, value in enumerate(product, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.border = thin_border
        if col_idx == 4:  # Price column
            cell.number_format = '#,##0.00'

# Auto-fit column widths (approximate)
column_widths = [12, 40, 18, 14, 80, 14, 22]
for i, width in enumerate(column_widths, 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

# Freeze the header row
ws.freeze_panes = "A2"

# Save
wb.save("d:/Code/Projects/datacrumbs/capstone_project_3/Company_products.xlsx")
print("✅ Company_products.xlsx created successfully with 15 products!")
